import glob
import math
import os
import re
from io import BytesIO

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Inches, Pt

TEMPLATE_DIR = os.path.join(os.path.dirname(__file__), "..", "assets", "templates")

BODY_FONT_PT = 14
CHARS_PER_LINE = 68
MAX_LINES_PER_SLIDE = 9


def _template_path():
    matches = sorted(glob.glob(os.path.join(TEMPLATE_DIR, "*.pptx")))
    return matches[0] if matches else None


def _clean(text):
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(text or "")).strip()


def _find_layouts(prs):
    """Pick layouts by placeholder signature because the template's layout names are opaque."""
    title_layout = content_layout = None
    best_width = 0
    for layout in prs.slide_layouts:
        types = [p.placeholder_format.type for p in layout.placeholders]
        if (
            title_layout is None
            and types.count(PP_PLACEHOLDER.TITLE) == 1
            and types.count(PP_PLACEHOLDER.SUBTITLE) == 2
            and PP_PLACEHOLDER.BODY not in types
        ):
            title_layout = layout
        if (
            PP_PLACEHOLDER.TITLE in types
            and PP_PLACEHOLDER.SUBTITLE in types
            and types.count(PP_PLACEHOLDER.BODY) == 1
            and PP_PLACEHOLDER.PICTURE not in types
        ):
            body = next(p for p in layout.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.BODY)
            if body.width > best_width:
                best_width, content_layout = body.width, layout
    layouts = list(prs.slide_layouts)
    return title_layout or layouts[0], content_layout or layouts[min(9, len(layouts) - 1)]


def _drop_sample_slides(prs):
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        prs.part.drop_rel(sld_id.rId)
        sld_id_lst.remove(sld_id)


def _remove_empty_placeholders(slide):
    for shape in list(slide.placeholders):
        if shape.has_text_frame and not shape.text_frame.text.strip():
            shape._element.getparent().remove(shape._element)
        elif shape.placeholder_format.type in (PP_PLACEHOLDER.PICTURE, PP_PLACEHOLDER.SLIDE_NUMBER):
            shape._element.getparent().remove(shape._element)


def _placeholder(slide, ph_type, nth=0):
    # The default python-pptx layouts use OBJECT for their content box; the PSU template uses BODY.
    accepted = (ph_type, PP_PLACEHOLDER.OBJECT) if ph_type == PP_PLACEHOLDER.BODY else (ph_type,)
    found = [p for p in slide.placeholders if p.placeholder_format.type in accepted]
    return found[nth] if len(found) > nth else None


def _set_text(shape, text):
    if shape is not None and text:
        shape.text_frame.text = _clean(text)


def _item_text(item):
    return item if isinstance(item, str) else f"{item[0]}: {item[1]}"


def _lines(item):
    return sum(max(1, math.ceil(len(part) / CHARS_PER_LINE)) for part in _item_text(item).split("\n"))


def _split_long(text):
    """Break one oversize string into word-boundary pieces that each fit on a slide."""
    limit = CHARS_PER_LINE * (MAX_LINES_PER_SLIDE - 1)
    pieces, current = [], ""
    for word in text.split(" "):
        if len(current) + len(word) + 1 > limit and current:
            pieces.append(current)
            current = word
        else:
            current = f"{current} {word}".strip()
    if current:
        pieces.append(current)
    return pieces


def _chunk(items):
    """Group items (str or (bold_term, text)) into slide-sized pages by estimated line count."""
    expanded = []
    for item in items:
        if _lines(item) > MAX_LINES_PER_SLIDE:
            if isinstance(item, str):
                expanded.extend(_split_long(item))
            else:
                expanded.extend((item[0], piece) for piece in _split_long(item[1]))
        else:
            expanded.append(item)

    pages, page, used = [], [], 0
    for item in expanded:
        n = _lines(item)
        if page and used + n > MAX_LINES_PER_SLIDE:
            pages.append(page)
            page, used = [], 0
        page.append(item)
        used += n
    if page:
        pages.append(page)
    return pages


def _fill_body(body, items):
    tf = body.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        if isinstance(item, str):
            run = p.add_run()
            run.text = _clean(item)
            run.font.size = Pt(BODY_FONT_PT)
        else:
            term, text = item
            bold = p.add_run()
            bold.text = f"{_clean(term)}: "
            bold.font.bold = True
            bold.font.size = Pt(BODY_FONT_PT)
            rest = p.add_run()
            rest.text = _clean(text)
            rest.font.size = Pt(BODY_FONT_PT)


def _add_content_slides(prs, layout, title, subtitle, items):
    pages = _chunk([i for i in items if i])
    for n, page in enumerate(pages):
        slide = prs.slides.add_slide(layout)
        suffix = f" ({n + 1}/{len(pages)})" if len(pages) > 1 else ""
        _set_text(slide.shapes.title, title + suffix)
        _set_text(_placeholder(slide, PP_PLACEHOLDER.SUBTITLE), subtitle)
        body = _placeholder(slide, PP_PLACEHOLDER.BODY)
        if body is not None:
            _fill_body(body, page)
        _remove_empty_placeholders(slide)


def create_practice_deck(package, course, include_questions=True, include_definitions=True, transcript=None):
    """Render a StudyPackage as a flashcard-style deck on the PSU template and return it as a BytesIO."""
    template = _template_path()
    if template:
        prs = Presentation(template)
        _drop_sample_slides(prs)
        title_layout, content_layout = _find_layouts(prs)
    else:
        # No template installed: fall back to python-pptx's default 16:9 title and title+content layouts.
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
        title_layout, content_layout = prs.slide_layouts[0], prs.slide_layouts[1]

    cover = prs.slides.add_slide(title_layout)
    _set_text(cover.shapes.title, f"{course} Study Package")
    _set_text(_placeholder(cover, PP_PLACEHOLDER.SUBTITLE, 0), "Review and practice")
    _set_text(_placeholder(cover, PP_PLACEHOLDER.SUBTITLE, 1), course)
    _remove_empty_placeholders(cover)

    _add_content_slides(prs, content_layout, "Session Summary", "What we worked through", [package.session_summary])

    if include_definitions and package.concepts:
        _add_content_slides(
            prs, content_layout, "Key Terms", "Say each definition in your own words",
            [(c.term, c.definition) for c in package.concepts],
        )

    if include_questions:
        for i, q in enumerate(package.practice_questions, 1):
            _add_content_slides(prs, content_layout, f"Practice Question {i}", "Try it before you flip the card", [q.question])
            _add_content_slides(
                prs, content_layout, f"Hints: Question {i}", "Guiding steps, not the answer",
                [("Hint", q.hint)] + [f"{n}. {step}" for n, step in enumerate(q.guidance_steps, 1)],
            )

    if package.pitfalls:
        _add_content_slides(prs, content_layout, "Watch Out For", "Common pitfalls from this session", list(package.pitfalls))

    if transcript:
        items = [(speaker, text) for speaker, text in transcript]
        _add_content_slides(prs, content_layout, "Appendix: Transcript", course, items)

    buffer = BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer
